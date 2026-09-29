"""Daily exception exports for the Excel review workbook."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

EXCEL_COLUMNS = [
    "exception_id",
    "mark_date",
    "book",
    "instrument",
    "asset_class",
    "fv_level",
    "check_name",
    "severity",
    "mark",
    "reference_price",
    "deviation_bps",
    "mv_impact_usd",
    "suspected_cause",
    "draft_commentary",
    "status",
]


def daily_exceptions_frame(repo, mark_date: date | None = None) -> pd.DataFrame:
    """Build the exact Excel review column set."""
    sql = """
    SELECT e.exception_id, e.mark_date, p.book, i.ticker AS instrument,
           i.asset_class, i.fv_level, e.check_name, e.severity, e.mark,
           e.reference_price, e.deviation_bps, e.mv_impact_usd,
           e.suspected_cause, c.draft_text AS draft_commentary, e.status
    FROM exceptions e
    JOIN positions p ON p.position_id=e.position_id
    JOIN instruments i ON i.instrument_id=p.instrument_id
    LEFT JOIN commentary c ON c.exception_id=e.exception_id
    """
    params: tuple = ()
    if mark_date is not None:
        sql += " WHERE e.mark_date=?"
        params = (mark_date.isoformat(),)
    sql += " ORDER BY e.severity, e.exception_id"
    frame = repo.query(sql, params)
    for column in EXCEL_COLUMNS:
        if column not in frame:
            frame[column] = None
    return frame[EXCEL_COLUMNS]


def export_daily_exceptions(
    repo,
    output_dir: str | Path,
    mark_date: date | None = None,
) -> tuple[Path, Path]:
    """Write daily exceptions as CSV and openpyxl-generated XLSX."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    if mark_date is None:
        latest = repo.query("SELECT MAX(mark_date) AS mark_date FROM exceptions")
        if not latest.empty and pd.notna(latest.iloc[0]["mark_date"]):
            mark_date = date.fromisoformat(str(latest.iloc[0]["mark_date"]))
    export_date = mark_date or date.today()
    frame = daily_exceptions_frame(repo, mark_date)
    csv_path = output / f"daily_exceptions_{export_date:%Y%m%d}.csv"
    xlsx_path = output / f"daily_exceptions_{export_date:%Y%m%d}.xlsx"
    frame.to_csv(csv_path, index=False)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    sheet.append(EXCEL_COLUMNS)
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if sheet.max_row >= 2:
        table = Table(displayName="ExceptionReview", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    workbook.save(xlsx_path)
    return csv_path, xlsx_path
